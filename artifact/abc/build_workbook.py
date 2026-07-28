"""Build the ABC-analysis workbook.

Values are written as literals, not formulas. LibreOffice is not functional in
this environment, so openpyxl formulas would ship with no cached values and read
back as empty in pandas / Claude / most previewers - useless for a file whose
whole purpose is to be read by an analysis tool. The source of truth is the
sales database plus artifact/build-data.ts; regenerate rather than hand-edit.
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

S = "/tmp/claude-0/-home-user-salesdata/4d497f07-3bf1-503f-9f9f-1477240e92c5/scratchpad"
D = json.load(open(f"{S}/abc-data.json"))
meta, months12 = D["meta"], D["months12"]

A_CUT, B_CUT = 0.80, 0.95

ARIAL = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
BODY = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=13, bold=True)
CLS_FONT = {"A": Font(name=ARIAL, size=10, bold=True, color="9C0006"),
            "B": Font(name=ARIAL, size=10, color="9C6500"),
            "C": Font(name=ARIAL, size=10, color="595959")}
CLS_FILL = {"A": PatternFill("solid", fgColor="FFC7CE"),
            "B": PatternFill("solid", fgColor="FFEB9C"),
            "C": PatternFill("solid", fgColor="EDEDED")}
BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))
MONEY, PCT, INT, DEC = '#,##0.00', '0.0%', '#,##0', '#,##0.0'

wb = Workbook()

# ---------------- README ----------------
ws = wb.active
ws.title = "README"
L = [
    ("Sales data extract for ABC analysis", TITLE),
    ("", BODY),
    (f"Generated {meta['generated']} from the Japan Stationery / Elite Camp sales database.", BODY),
    (f"Full history: {meta['allMonths'][0]} to {meta['allMonths'][1]} ({meta['rowsAll']:,} order line items).", BODY),
    (f"Rolling 12 months: {meta['window12'][0]} to {meta['window12'][1]} ({meta['rows12']:,} order line items).", BODY),
    (f"Month {meta['partialExcluded']} is still in progress and is EXCLUDED from every sheet.", BODY),
    ("", BODY),
    ("Sheets", BOLD),
    ("ABC JS 12M          Japan Stationery, rolling 12 months, ranked by revenue. START HERE.", BODY),
    ("ABC EC 12M          Elite Camp, rolling 12 months, ranked by revenue.", BODY),
    ("ABC JS All          Japan Stationery, full history.", BODY),
    ("ABC EC All          Elite Camp, full history.", BODY),
    ("Monthly Units 12M   Units per product per month, for demand variability / XYZ analysis.", BODY),
    ("", BODY),
    ("Method", BOLD),
    ("Products are ranked by revenue within each shop. Cumulative % of that shop's revenue is", BODY),
    ("accumulated down the ranking. A product is class A while cumulative % is at or below", BODY),
    (f"{A_CUT:.0%}, class B up to {B_CUT:.0%}, class C thereafter. This is ABC by REVENUE CONTRIBUTION.", BODY),
    ("The Cumulative % column is included on every sheet, so any other cut-off can be applied", BODY),
    ("without recomputing anything - just re-read that column.", BODY),
    ("", BODY),
    ("Values are stored as numbers, not formulas, so every tool reads them directly.", BODY),
    ("To refresh after importing new sales files, regenerate from the database rather than", BODY),
    ("editing cells by hand.", BODY),
    ("", BODY),
    ("What this data CANNOT tell you", BOLD),
    ("1. NO COST OR MARGIN. The source system holds no unit cost, supplier price or margin.", BODY),
    ("   So this is ABC by revenue, NOT by cost of consumption and NOT by profit. A", BODY),
    ("   high-revenue, thin-margin line looks more important here than it really is. If you", BODY),
    ("   can supply a cost per product, a proper margin-weighted ABC becomes possible.", BODY),
    ("2. NO STOCK DATA. No stock on hand, lead time, MOQ or reorder point, so safety stock", BODY),
    ("   and reorder points cannot be derived from this file alone.", BODY),
    ("3. Products are grouped by a cleaned-up product name, because the same item gets", BODY),
    ("   relisted under different marketing titles. Name Variants shows how many raw listing", BODY),
    ("   titles merged into each row - review the high ones before trusting the ranking.", BODY),
    ("4. Only Completed (Shopee) and confirmed (Lazada) orders count; cancellations and", BODY),
    ("   returns are already excluded.", BODY),
    ("5. Each Lazada line is one unit, so Lazada unit counts are reliable, but its per-unit", BODY),
    ("   price is derived from the paid amount.", BODY),
    ("6. Olight sells through BOTH shops. Rows stay with the shop that made the sale, so if", BODY),
    ("   the shops share Olight stock, combine both shops before ordering.", BODY),
    ("7. Shopee and Lazada revenue are split per product, but stock is presumably shared -", BODY),
    ("   add the two for purchasing decisions.", BODY),
    ("8. Revenue is product subtotal after discounts; it excludes shipping and platform fees.", BODY),
]
for i, (t, f) in enumerate(L, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(vertical="top")
ws.column_dimensions["A"].width = 98

HEAD = ["Rank", "Product", "Brand", "Units", "Revenue (RM)", "Avg Unit Price (RM)",
        "Shopee Rev (RM)", "Lazada Rev (RM)", "Order Lines", "Months Active",
        "Name Variants", "First Sale", "Last Sale", "% of Revenue", "Cumulative %", "Class"]
WIDTHS = [6, 52, 12, 9, 14, 17, 14, 14, 11, 12, 12, 11, 11, 12, 13, 7]


def abc_sheet(title, items, period):
    s = wb.create_sheet(title)
    s["A1"] = f"{title} - ranked by revenue - {period}"
    s["A1"].font = TITLE
    s["A2"] = f"Class A to {A_CUT:.0%} cumulative, B to {B_CUT:.0%}, C after. ABC by revenue; no cost data exists."
    s["A2"].font = BODY

    hdr = 4
    for j, h in enumerate(HEAD, start=1):
        c = s.cell(row=hdr, column=j, value=h)
        c.font = H_FONT
        c.fill = H_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    total_rev = sum(x["revenue"] for x in items)
    total_units = sum(x["units"] for x in items)
    total_lines = sum(x["lines"] for x in items)
    running = 0.0
    counts = {"A": 0, "B": 0, "C": 0}
    cls_rev = {"A": 0.0, "B": 0.0, "C": 0.0}
    first = hdr + 1

    for i, it in enumerate(items):
        r = first + i
        running += it["revenue"]
        share = it["revenue"] / total_rev if total_rev else 0.0
        cum = running / total_rev if total_rev else 0.0
        cls = "A" if cum <= A_CUT else ("B" if cum <= B_CUT else "C")
        counts[cls] += 1
        cls_rev[cls] += it["revenue"]
        vals = [i + 1, it["name"], it["brand"], it["units"], round(it["revenue"], 2),
                round(it["revenue"] / it["units"], 2) if it["units"] else 0,
                round(it["shopee"], 2), round(it["lazada"], 2), it["lines"],
                it["monthsActive"], it["variants"], it["first"], it["last"],
                round(share, 6), round(cum, 6), cls]
        fmts = [INT, None, None, INT, MONEY, MONEY, MONEY, MONEY, INT, INT, INT,
                None, None, PCT, PCT, None]
        for j, (v, fmt) in enumerate(zip(vals, fmts), start=1):
            c = s.cell(row=r, column=j, value=v)
            c.font = BODY
            c.border = BORDER
            if fmt:
                c.number_format = fmt
        cc = s.cell(row=r, column=16)
        cc.font = CLS_FONT[cls]
        cc.fill = CLS_FILL[cls]
        cc.alignment = Alignment(horizontal="center")

    last = first + len(items) - 1
    t = last + 2
    s.cell(row=t, column=2, value="TOTAL").font = BOLD
    for j, v, fmt in ((4, total_units, INT), (5, round(total_rev, 2), MONEY),
                      (9, total_lines, INT)):
        c = s.cell(row=t, column=j, value=v)
        c.font = BOLD
        c.number_format = fmt

    b = t + 2
    s.cell(row=b, column=2, value="Class summary").font = BOLD
    for j, h in ((4, "Products"), (5, "Revenue (RM)"), (14, "% of Revenue")):
        s.cell(row=b, column=j, value=h).font = BOLD
    for k, cls in enumerate(("A", "B", "C")):
        rr = b + 1 + k
        c = s.cell(row=rr, column=2, value=cls)
        c.font = CLS_FONT[cls]
        c.fill = CLS_FILL[cls]
        c.alignment = Alignment(horizontal="center")
        s.cell(row=rr, column=4, value=counts[cls]).number_format = INT
        s.cell(row=rr, column=5, value=round(cls_rev[cls], 2)).number_format = MONEY
        s.cell(row=rr, column=14,
               value=round(cls_rev[cls] / total_rev if total_rev else 0, 6)).number_format = PCT
        for j in (4, 5, 14):
            s.cell(row=rr, column=j).font = BODY
    s.cell(row=b + 5, column=2,
           value=f"Source: sales database, {period}. Revenue is product subtotal after discounts, "
                 f"excluding shipping and platform fees.").font = BODY

    for j, w in enumerate(WIDTHS, start=1):
        s.column_dimensions[get_column_letter(j)].width = w
    s.freeze_panes = f"C{first}"
    return counts, cls_rev, total_rev


p12 = f"{meta['window12'][0]} to {meta['window12'][1]}"
pall = f"{meta['allMonths'][0]} to {meta['allMonths'][1]}"
summary = {}
summary["JS 12M"] = abc_sheet("ABC JS 12M", [x for x in D["a12"] if x["shop"] == "Japan Stationery"], p12)
summary["EC 12M"] = abc_sheet("ABC EC 12M", [x for x in D["a12"] if x["shop"] == "Elite Camp"], p12)
summary["JS All"] = abc_sheet("ABC JS All", [x for x in D["aAll"] if x["shop"] == "Japan Stationery"], pall)
summary["EC All"] = abc_sheet("ABC EC All", [x for x in D["aAll"] if x["shop"] == "Elite Camp"], pall)

# ---------------- Monthly units ----------------
mu = wb.create_sheet("Monthly Units 12M")
mu["A1"] = f"Units sold per product per month - {p12}"
mu["A1"].font = TITLE
mu["A2"] = "Use for demand variability (XYZ): steady sellers vs spiky ones. Zero means no sales that month."
mu["A2"].font = BODY
hdr = 4
head = ["Shop", "Product", "Brand"] + months12 + ["Total Units", "Months w/ Sales", "Avg / Active Month"]
for j, h in enumerate(head, start=1):
    c = mu.cell(row=hdr, column=j, value=h)
    c.font = H_FONT
    c.fill = H_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
rows_sorted = sorted(D["monthly"], key=lambda x: (x["shop"], -sum(x["u"].values())))
nm = len(months12)
for i, it in enumerate(rows_sorted):
    r = hdr + 1 + i
    mu.cell(row=r, column=1, value=it["shop"])
    mu.cell(row=r, column=2, value=it["name"])
    mu.cell(row=r, column=3, value=it["brand"])
    tot = 0
    active = 0
    for k, m in enumerate(months12):
        u = it["u"].get(m, 0)
        tot += u
        active += 1 if u > 0 else 0
        mu.cell(row=r, column=4 + k, value=u).number_format = INT
    tc = 4 + nm
    mu.cell(row=r, column=tc, value=tot).number_format = INT
    mu.cell(row=r, column=tc + 1, value=active).number_format = INT
    mu.cell(row=r, column=tc + 2,
            value=round(tot / active, 1) if active else 0).number_format = DEC
    for j in range(1, tc + 3):
        mu.cell(row=r, column=j).font = BODY
        mu.cell(row=r, column=j).border = BORDER
mu.column_dimensions["A"].width = 18
mu.column_dimensions["B"].width = 52
mu.column_dimensions["C"].width = 12
for k in range(nm):
    mu.column_dimensions[get_column_letter(4 + k)].width = 9
for k in range(3):
    mu.column_dimensions[get_column_letter(4 + nm + k)].width = 15
mu.freeze_panes = f"D{hdr+1}"

out = "/home/user/salesdata/artifact/abc-analysis-data.xlsx"
wb.save(out)
print("saved", out)
for k, (counts, cls_rev, tot) in summary.items():
    parts = " ".join(f"{c}={counts[c]}({cls_rev[c]/tot:.1%})" for c in "ABC")
    print(f"  {k}: total RM {tot:,.2f}  {parts}")
